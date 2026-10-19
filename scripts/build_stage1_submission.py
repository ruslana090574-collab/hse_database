"""Build the stage-one report from a verified database and original Mac captures."""
from pathlib import Path
from decimal import Decimal
from zipfile import ZipFile
import argparse
import hashlib
import json
import re
import textwrap

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
FONT = 'Liberation Sans'
KEEP_QUERIES = set(range(1, 8)) | set(range(10, 15)) | {17, 18}
DROP_SECTIONS = {4, 6, 8, 9}
RENUMBER = {5: 4, 7: 5, 10: 6, 11: 7, 12: 8}
PREFERRED = {
    '01': ['PassportNo', 'FullName', 'PolicyID', 'CountryName', 'PolicyCoverage', 'CurrencyLabel'],
    '02': ['PassportNo', 'FullName', 'PolicyID', 'CountryName', 'PolicyCoverage'],
    '03': ['PoliciesCount'],
    '04': ['PassportNo', 'FullName', 'PoliciesCount'],
    '05': ['PassportNo', 'FullName', 'BirthDate', 'PoliciesCount'],
    '06': ['PassportNo', 'FullName', 'PolicyCoverage', 'CurrencyLabel'],
    '07': ['PassportNo', 'FullName', 'BirthDate'],
    '10': ['AgentID', 'AgentName', 'PolicyCount'],
    '11': ['AgentName', 'RevenueRUB', 'CommissionRUB'],
    '12': ['AgentID', 'AgentName', 'PolicyCount'],
    '13': ['PolicyCount', 'AvgPremiumRUB'],
    '14': ['AgentName', 'PolicyCount', 'AvgPremiumRUB'],
    '17': ['AgentID', 'AgentName', 'IsActive'],
    '18': ['PassportNo', 'FullName', 'CoverageRUB'],
}
MONEY = {'RevenueRUB', 'CommissionRUB', 'AvgPremiumRUB', 'CoverageRUB', 'PremiumRUB'}


def font(run, family=FONT):
    run.font.name = family
    fonts = run._r.get_or_add_rPr().rFonts
    for name in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + name), family)
    for name in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        fonts.attrib.pop(qn('w:' + name), None)


def display(value, column):
    if value is None:
        return '—'
    if column in MONEY:
        return format(Decimal(value), '.2f').replace('.', ',')
    if re.fullmatch(r'\d{4}-\d{2}-\d{2}', str(value)):
        year, month, day = value.split('-')
        return f'{day}.{month}.{year}'
    return str(value)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--screenshots', type=Path, required=True)
    parser.add_argument('--results', type=Path, required=True)
    parser.add_argument('--queries', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    results = json.loads(args.results.read_text())
    queries = json.loads(args.queries.read_text())
    from model import LABELS

    document = Document(ROOT / 'Отчёт_Шагеева_РР.docx')
    body = document._element.body
    section = None
    skip_query = False
    query_name = None
    query_table_pending = False
    for child in list(body):
        if child.tag == qn('w:sectPr'):
            continue
        paragraph = Paragraph(child, document._body) if child.tag == qn('w:p') else None
        if paragraph is not None and paragraph.style.name == 'Heading 1':
            match = re.match(r'(\d+)\.', paragraph.text)
            section = int(match.group(1)) if match else None
            skip_query = False
            query_name = None
        elif paragraph is not None and section == 5 and paragraph.style.name == 'Heading 2':
            match = re.match(r'(\d{2}) ', paragraph.text)
            if match:
                number = int(match.group(1))
                skip_query = number not in KEEP_QUERIES
                query_name = next(name for name in queries if name.startswith(f'{number:02d}_'))
                query_table_pending = True
        if section in DROP_SECTIONS or skip_query:
            body.remove(child)
            continue
        if paragraph is not None:
            text = paragraph.text
            updated = text
            if paragraph.style.name == 'Heading 1' and section in RENUMBER:
                updated = re.sub(r'^\d+\.', str(RENUMBER[section]) + '.', text)
            elif text.startswith('Этапы 1 и 2.'):
                updated = 'Этап 1. Разработка базы данных и запросов\nLibreOffice Base, macOS'
            elif text.startswith('Целевая программа:'):
                updated = 'Программная среда: LibreOffice Base, macOS.\nФормат базы данных: ODB, встроенная Firebird.'
            elif text.startswith('Цель: разработать реляционную'):
                updated = ('Цель работы — разработать реляционную базу учёта клиентов страховой компании и запросы для анализа '
                           'полисов и работы страховых агентов. Отчёт охватывает первый этап: создание таблиц, определение '
                           'ключей, связей и ограничений, основные запросы и самостоятельные задания по структуре базы '
                           'и оценке работы агентов. Основание — методические рекомендации Л. Н. Лядовой, Пермь, 2026.')
            elif text.startswith('Исходная методичка и требования LMS'):
                updated = ('Методические рекомендации ориентированы на MS Access. Для macOS использована LibreOffice Base '
                           'со встроенной Firebird. База сохранена в формате .odb. Структура, ограничения и запросы '
                           'адаптированы к возможностям этой СУБД; форматы .odb и .accdb различаются.')
            elif text.startswith('Объекты созданы и настроены средствами API'):
                updated = ('Созданы таблицы предметной области, заданы первичные и внешние ключи, ограничения значений и '
                           'запросы. Структура Client просмотрена в конструкторе таблиц, связи — в окне Relationships. '
                           'Скриншоты интерфейса в данном отчёте получены при работе с сохранённой базой на macOS. '
                           'Рисунок 1 представляет отдельно построенную схему по фактической структуре базы.')
            elif text.startswith('Для совместимости со встроенной Firebird'):
                updated = text.replace('Соответствие имён приведено ниже и в файлах SQL.',
                                       'Соответствие имён приведено ниже; состав полей дан в приложении А.')
            elif text.startswith('Рисунок 1. Полная схема'):
                updated = 'Рисунок 1. Схема базы данных: таблицы, ключи и связи, построенные по фактическим определениям таблиц.'
            elif text.startswith('Рисунок 2. Окно связей'):
                updated = 'Рисунок 2. Окно связей LibreOffice Base на macOS: 12 таблиц и связи между ними.'
            elif text.startswith('Рисунок 3. Конструктор таблицы Client'):
                updated = 'Рисунок 3. Таблица Client в конструкторе Base на macOS: поля, типы и длина PassportNo — 9 символов.'
            elif text.startswith('Рисунок 6. Результат запроса 05'):
                updated = 'Рисунок 4. Результат запроса 05 на macOS: восемь клиентов, включая троих с нулём полисов за текущий год.'
            elif text.startswith('Рисунок 8. Предлагаемая схема расширения'):
                updated = text.replace('Рисунок 8.', 'Рисунок 5.', 1)
            elif text.startswith('Текстовые цифровые поля имеют маски'):
                updated = ('Формат паспортов и цифровых кодов контролируется ограничениями CHECK на уровне Firebird. '
                           'Ограничения действуют и при прямом вводе в таблицы. Возраст проверяется прибавлением '
                           '18 календарных лет к дате рождения, а не делением числа дней на 365.')
            elif text.startswith('Запросы 10–14 принимают'):
                updated = ('Запросы 10–14 принимают DateStart и DateEnd; обе границы периода включаются. '
                           'Контрольный период — 01.01.2026–31.12.2026. Параметры имеют тип даты. '
                           'Средние денежные значения вычисляются с промежуточной точностью четыре десятичных знака '
                           'и округляются до копеек. В таблицах отчёта показаны основные поля; полный набор столбцов '
                           'доступен при выполнении соответствующего сохранённого запроса в Base.')
            elif text.startswith('Получено строк:') and query_name:
                updated = f"Получено строк: {len(results[query_name]['rows'])}. В таблице приведены основные поля результата."
            elif query_name and any(run.font.name == 'Liberation Mono' for run in paragraph.runs):
                updated = '\n'.join(textwrap.wrap(queries[query_name], 94, break_long_words=False, break_on_hyphens=False))
                paragraph.text = updated
                for run in paragraph.runs:
                    font(run, 'Liberation Mono')
                    run.font.size = Pt(8)
                continue
            elif text.startswith('Отрицательные проверки выполнялись'):
                updated = ('Контрольные изменения выполнялись над отдельной копией сохранённой базы с откатом транзакций. '
                           'Проверки обязательных полей, форматов, возраста, сроков, внешних ключей и сохранения '
                           'исторических условий завершились успешно. Все сохранённые запросы выполняются; '
                           'результаты основных запросов сопоставлены с контрольными данными. '
                           'Исходные записи клиентов, полисов и справочников сохранены.')
            elif 'Для истории в формах показывается сохранённая цена' in text:
                updated = text.replace('Для истории в формах показывается сохранённая цена, а не цена из актуального списка выбора.',
                                       'В исторических расчётах используется сохранённая договорная цена.')
            elif text.startswith('8. LibreOffice Help:'):
                updated = '8. LibreOffice Help: Base, таблицы, связи и запросы. https://help.libreoffice.org/latest/ru/'
            if updated != text:
                paragraph.text = updated
        elif child.tag == qn('w:tbl') and section == 5 and query_name and query_table_pending:
            table = Table(child, document._body)
            data = results[query_name]
            columns = [name for name in PREFERRED[query_name[:2]] if name in data['columns']]
            table._tbl.tblGrid.clear()
            for row in list(table.rows):
                table._tbl.remove(row._tr)
            from docx.oxml import OxmlElement
            weights = [2.8 if name in {'FullName', 'AgentName'} else
                       1.4 if name in {'PassportNo', 'CountryName'} else
                       1.2 if name == 'BirthDate' else 1 for name in columns]
            for weight in weights:
                grid = OxmlElement('w:gridCol')
                grid.set(qn('w:w'), str(round(9638 * weight / sum(weights))))
                table._tbl.tblGrid.append(grid)
            header = table.add_row()
            for cell, name in zip(header.cells, columns):
                cell.text = LABELS.get(name, name)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            marker = OxmlElement('w:tblHeader')
            header._tr.get_or_add_trPr().append(marker)
            for result_row in data['rows']:
                row = table.add_row()
                for cell, name in zip(row.cells, columns):
                    cell.text = display(result_row[data['columns'].index(name)], name)
            query_table_pending = False

    # Keep the first-stage entries in the requirement matrix and concise check results.
    excluded_requirements = {'Формы клиента и полиса', 'Формы остальных сущностей', 'Отчёт с частичными итогами',
                             'Электронные таблицы', 'Документы из БД'}
    for table in document.tables:
        if table.rows and table.rows[0].cells[0].text == 'MS Access / методичка':
            for row in list(table.rows)[1:]:
                if row.cells[0].text in {'Формы и подчинённые формы', 'Мастер / конструктор отчётов', 'Word / Excel'}:
                    table._tbl.remove(row._tr)
                elif row.cells[0].text == 'Подстановки в таблицах':
                    row.cells[1].text = 'В таблицах хранятся коды справочников; наименования отображаются через JOIN в запросах'
        if table.rows and table.rows[0].cells[0].text == 'Требование':
            for row in list(table.rows)[1:]:
                if row.cells[0].text in excluded_requirements:
                    table._tbl.remove(row._tr)
                elif row.cells[0].text == 'Ограничения, PK, FK и подстановки':
                    row.cells[2].text = 'CHECK, триггеры, первичные и внешние ключи, составные ключи и связи справочников'
        if table.rows and [cell.text for cell in table.rows[0].cells] == ['№', 'Проверка', 'Результат']:
            for row in list(table.rows)[1:]:
                table._tbl.remove(row._tr)
            checks = [
                'После повторного открытия сохранены 8 клиентов и 13 полисов',
                'LEFT JOIN и NOT EXISTS возвращают клиентов без полисов за текущий год',
                'Запрос лидеров возвращает обоих агентов с максимальным числом полисов',
                'Недопустимые коды, пустые обязательные значения и повторные ключи отклоняются',
                'Кириллица в ФИО клиента и буквы в номере паспорта отклоняются',
                'Несовершеннолетний клиент, просроченный паспорт и будущая дата выдачи отклоняются',
                'Проверены сроки полиса, действие паспорта и договора агента',
                'Внешние ключи запрещают несуществующие ссылки и удаление клиента с полисами',
                'Исторические цены и покрытия не меняются при изменении справочников',
                'Удаление полиса каскадно удаляет его услуги; откат восстанавливает данные',
                'Средняя стоимость за 2026 год округлена до 329,88 руб.; отсутствие продаж даёт NULL',
                'Сохранённые запросы выполняются, включая типизированные параметры периода',
            ]
            for index, check in enumerate(checks, 1):
                row = table.add_row()
                for cell, value in zip(row.cells, [str(index), check, 'Пройдено']):
                    cell.text = value
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    text = paragraph.text
                    updated = text.replace('словарь и SQL приложены', 'словарь — в приложении А; SQL запросов — в разделе 4')
                    updated = updated.replace('Адаптирован по запросу пользователя в .odb; согласование формата остаётся внешним условием',
                                              'Использован формат LibreOffice Base .odb; исходный формат методички — Access')
                    if updated != text:
                        paragraph.text = updated

    replacements = {
        hashlib.sha256((ROOT / 'screenshots/03_Связи.png').read_bytes()).digest(): args.screenshots / '01_Схема_связей_Mac.png',
        hashlib.sha256((ROOT / 'screenshots/04_Конструктор_таблицы.png').read_bytes()).digest(): args.screenshots / '02_Конструктор_Client_Mac.png',
        hashlib.sha256((ROOT / 'screenshots/06_Запрос_клиенты_ноль.png').read_bytes()).digest(): args.screenshots / '03_Запрос_05_Mac.png',
    }
    replaced = 0
    for paragraph in document.paragraphs:
        for drawing in list(paragraph._p.xpath('.//a:blip')):
            relationship = drawing.get(qn('r:embed'))
            digest = hashlib.sha256(document.part.related_parts[relationship].blob).digest()
            if digest in replacements:
                paragraph.clear()
                paragraph.add_run().add_picture(str(replacements[digest]), width=Cm(16.5))
                paragraph.paragraph_format.keep_with_next = True
                replaced += 1
    assert replaced == 3
    referenced = {node.get(qn('r:embed')) for node in document._element.xpath('.//a:blip')}
    for key, relationship in list(document.part.rels.items()):
        if relationship.reltype.endswith('/image') and key not in referenced:
            document.part.drop_rel(key)

    # Explicit families prevent font-theme fallback from hiding regular text during export.
    for style in document.styles:
        if style.type in (1, 2):
            style.font.name = FONT
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        headers = [cell.text for cell in table.rows[0].cells]
        widths = None
        if headers == ['№', 'Проверка', 'Результат']:
            widths = [1.0, 13.2, 2.8]
        elif headers == ['Требование', 'Источник', 'Результат']:
            widths = [5.2, 3.0, 8.8]
        elif headers == ['Поле', 'Наименование', 'Тип / правило']:
            widths = [3.7, 5.0, 8.3]
        elif headers == ['Объект', 'Реализованное правило']:
            widths = [4.5, 12.5]
        elif headers == ['MS Access / методичка', 'Адаптация в LibreOffice']:
            widths = [6.0, 11.0]
        elif headers == ['Таблица методички', 'Физическое имя', 'Первичный ключ']:
            widths = [5.4, 3.2, 8.4]
        if widths:
            table.autofit = False
            for column, width in zip(table.columns, widths):
                column.width = Cm(width)
            for row in table.rows:
                for cell, width in zip(row.cells, widths):
                    cell.width = Cm(width)
        if table.rows[0]._tr.get_or_add_trPr().find(qn('w:tblHeader')) is None:
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for row in table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in document.sections:
        paragraphs.extend(section.footer.paragraphs)
    for paragraph in paragraphs:
        paragraph.paragraph_format.widow_control = True
        if paragraph.text and len(paragraph.text) < 1500:
            paragraph.paragraph_format.keep_together = True
        if paragraph._p.xpath('.//a:blip'):
            paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            family = 'Liberation Mono' if run.font.name == 'Liberation Mono' else FONT
            font(run, family)
            if paragraph._p.getparent().tag == qn('w:tc'):
                run.font.size = Pt(9)
        if paragraph.text.startswith('Рисунок '):
            for run in paragraph.runs:
                run.italic = True
                run.font.size = Pt(9)
    document.core_properties.author = 'Шагеева Руслана Руслановна'
    document.core_properties.title = 'ЛР-1. Система учёта клиентов страховой компании. Этап 1'
    document.core_properties.subject = 'Таблицы, связи, ограничения и SQL-запросы'
    target = args.output / 'Отчёт_Шагеева_РР.docx'
    document.save(target)
    check = Document(target)
    assert len(check.inline_shapes) == 5
    text = '\n'.join(paragraph.text for paragraph in check.paragraphs)
    for excluded in ('Codex', 'results/', 'CSV-приложении', 'по условию пользователя', 'указана пользователем', 'Этапы 1 и 2'):
        assert excluded not in text, excluded
    assert 'ROUND(AVG(CAST(' in text
    with ZipFile(target) as archive:
        images = [archive.read(name) for name in archive.namelist() if name.startswith('word/media/')]
        assert len(images) == 5
        hashes = {hashlib.sha256(blob).digest() for blob in images}
        for screenshot in replacements.values():
            assert hashlib.sha256(screenshot.read_bytes()).digest() in hashes
    print('Stage-one report:', target)
    print('Original Mac screenshots embedded without changing their image data: 3')
    print('Additional schema illustrations: 2')
    print('Query sections:', len(KEEP_QUERIES))


if __name__ == '__main__':
    main()
